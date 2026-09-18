from decimal import Decimal

from django.http import HttpResponse

from openpyxl import Workbook
from openpyxl.styles import (
    Font,
    PatternFill,
    Alignment,
)
from openpyxl.utils import get_column_letter

from transactions.models import (
    SalesInvoice,
    SalesInvoiceItem,
)

from purchases.models import (
    PurchaseInvoice,
    PurchaseInvoiceItem,
)

from masters.models import (
    Product,
    Customer,
    Supplier,
)

from payments.models import (
    CustomerPayment,
    SupplierPayment,
)


# ==========================================================
# STYLES
# ==========================================================

HEADER_FILL = PatternFill(
    fill_type="solid",
    fgColor="1F4E78",
)

HEADER_FONT = Font(
    bold=True,
    color="FFFFFF",
)

TITLE_FONT = Font(
    bold=True,
    size=16,
)

BOLD_FONT = Font(
    bold=True,
)

CENTER = Alignment(
    horizontal="center",
    vertical="center",
)

LEFT = Alignment(
    horizontal="left",
)

RIGHT = Alignment(
    horizontal="right",
)


# ==========================================================
# COMMON FUNCTIONS
# ==========================================================

def apply_header_style(sheet, row):

    for cell in sheet[row]:

        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER


def auto_width(sheet):

    for column in sheet.columns:

        length = 0

        letter = get_column_letter(
            column[0].column
        )

        for cell in column:

            try:

                if len(str(cell.value)) > length:

                    length = len(str(cell.value))

            except Exception:
                pass

        sheet.column_dimensions[
            letter
        ].width = length + 5


def money(value):

    if value is None:
        value = Decimal("0.00")

    return float(value)


# ==========================================================
# RESPONSE
# ==========================================================

def workbook_response(workbook, filename):

    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        )
    )

    response[
        "Content-Disposition"
    ] = (
        f'attachment; filename="{filename}"'
    )

    workbook.save(response)

    return response

# ==========================================================
# SALES REPORT
# ==========================================================

def export_sales_excel():

    workbook = Workbook()

    sheet = workbook.active
    sheet.title = "Sales Report"

    # ------------------------------------------------------
    # TITLE
    # ------------------------------------------------------

    sheet["A1"] = "SALES REPORT"
    sheet["A1"].font = TITLE_FONT

    # ------------------------------------------------------
    # HEADERS
    # ------------------------------------------------------

    headers = [

        "Invoice No",
        "Invoice Date",
        "Customer",
        "Payment Type",

        "Subtotal",
        "Discount",
        "GST %",
        "GST Amount",
        "Grand Total",

    ]

    for col, header in enumerate(headers, start=1):

        cell = sheet.cell(
            row=3,
            column=col,
        )

        cell.value = header

    apply_header_style(sheet, 3)

    # ------------------------------------------------------
    # DATA
    # ------------------------------------------------------

    row = 4

    invoices = (
        SalesInvoice.objects
        .select_related("customer")
        .order_by("-invoice_number")
    )

    total_sales = Decimal("0.00")
    total_gst = Decimal("0.00")

    for invoice in invoices:

        sheet.cell(
            row=row,
            column=1,
            value=invoice.display_invoice_number,
        )

        sheet.cell(
            row=row,
            column=2,
            value=invoice.invoice_date.strftime(
                "%d-%m-%Y"
            ),
        )

        sheet.cell(
            row=row,
            column=3,
            value=invoice.customer.name,
        )

        sheet.cell(
            row=row,
            column=4,
            value=invoice.payment_type,
        )

        sheet.cell(
            row=row,
            column=5,
            value=money(invoice.subtotal),
        )

        sheet.cell(
            row=row,
            column=6,
            value=money(invoice.discount),
        )

        sheet.cell(
            row=row,
            column=7,
            value=float(invoice.gst_percent),
        )

        sheet.cell(
            row=row,
            column=8,
            value=money(invoice.gst_amount),
        )

        sheet.cell(
            row=row,
            column=9,
            value=money(invoice.grand_total),
        )

        total_sales += invoice.grand_total
        total_gst += invoice.gst_amount

        row += 1

    # ------------------------------------------------------
    # TOTALS
    # ------------------------------------------------------

    sheet.cell(
        row=row + 1,
        column=8,
        value="TOTAL",
    ).font = BOLD_FONT

    sheet.cell(
        row=row + 1,
        column=9,
        value=money(total_sales),
    ).font = BOLD_FONT

    sheet.cell(
        row=row + 2,
        column=8,
        value="TOTAL GST",
    ).font = BOLD_FONT

    sheet.cell(
        row=row + 2,
        column=9,
        value=money(total_gst),
    ).font = BOLD_FONT

    # ------------------------------------------------------
    # COLUMN WIDTH
    # ------------------------------------------------------

    auto_width(sheet)

    return workbook_response(
        workbook,
        "Sales_Report.xlsx",
    )

# ==========================================================
# PURCHASE REPORT
# ==========================================================

def export_purchase_excel():

    workbook = Workbook()

    sheet = workbook.active
    sheet.title = "Purchase Report"

    # ------------------------------------------------------
    # TITLE
    # ------------------------------------------------------

    sheet["A1"] = "PURCHASE REPORT"
    sheet["A1"].font = TITLE_FONT

    # ------------------------------------------------------
    # HEADERS
    # ------------------------------------------------------

    headers = [

        "Invoice No",
        "Invoice Date",
        "Supplier",
        "Payment Type",

        "Subtotal",
        "Discount",
        "GST %",
        "GST Amount",
        "Transport",
        "Grand Total",

    ]

    for col, header in enumerate(headers, start=1):

        cell = sheet.cell(
            row=3,
            column=col,
        )

        cell.value = header

    apply_header_style(sheet, 3)

    # ------------------------------------------------------
    # DATA
    # ------------------------------------------------------

    row = 4

    invoices = (
        PurchaseInvoice.objects
        .select_related("supplier")
        .order_by("-invoice_number")
    )

    total_purchase = Decimal("0.00")
    total_gst = Decimal("0.00")
    total_transport = Decimal("0.00")

    for invoice in invoices:

        sheet.cell(
            row=row,
            column=1,
            value=invoice.display_invoice_number,
        )

        sheet.cell(
            row=row,
            column=2,
            value=invoice.invoice_date.strftime(
                "%d-%m-%Y"
            ),
        )

        sheet.cell(
            row=row,
            column=3,
            value=invoice.supplier.name,
        )

        sheet.cell(
            row=row,
            column=4,
            value=invoice.payment_type,
        )

        sheet.cell(
            row=row,
            column=5,
            value=money(invoice.subtotal),
        )

        sheet.cell(
            row=row,
            column=6,
            value=money(invoice.discount),
        )

        sheet.cell(
            row=row,
            column=7,
            value=float(invoice.gst_percent),
        )

        sheet.cell(
            row=row,
            column=8,
            value=money(invoice.gst_amount),
        )

        sheet.cell(
            row=row,
            column=9,
            value=money(invoice.transport_charge),
        )

        sheet.cell(
            row=row,
            column=10,
            value=money(invoice.grand_total),
        )

        total_purchase += invoice.grand_total
        total_gst += invoice.gst_amount
        total_transport += invoice.transport_charge

        row += 1

    # ------------------------------------------------------
    # TOTALS
    # ------------------------------------------------------

    sheet.cell(
        row=row + 1,
        column=9,
        value="TOTAL",
    ).font = BOLD_FONT

    sheet.cell(
        row=row + 1,
        column=10,
        value=money(total_purchase),
    ).font = BOLD_FONT

    sheet.cell(
        row=row + 2,
        column=9,
        value="TOTAL GST",
    ).font = BOLD_FONT

    sheet.cell(
        row=row + 2,
        column=10,
        value=money(total_gst),
    ).font = BOLD_FONT

    sheet.cell(
        row=row + 3,
        column=9,
        value="TRANSPORT",
    ).font = BOLD_FONT

    sheet.cell(
        row=row + 3,
        column=10,
        value=money(total_transport),
    ).font = BOLD_FONT

    # ------------------------------------------------------
    # AUTO WIDTH
    # ------------------------------------------------------

    auto_width(sheet)

    return workbook_response(
        workbook,
        "Purchase_Report.xlsx",
    )

# ==========================================================
# STOCK REPORT
# ==========================================================

def export_stock_excel():

    workbook = Workbook()

    sheet = workbook.active
    sheet.title = "Stock Report"

    # ------------------------------------------------------
    # TITLE
    # ------------------------------------------------------

    sheet["A1"] = "STOCK REPORT"
    sheet["A1"].font = TITLE_FONT

    # ------------------------------------------------------
    # HEADERS
    # ------------------------------------------------------

    headers = [

        "Product",

        "Category",

        "Purchase Price",

        "Selling Price",

        "Current Stock",

        "Unit",

        "Stock Value",

        "Status",

    ]

    for col, header in enumerate(
        headers,
        start=1,
    ):

        cell = sheet.cell(
            row=3,
            column=col,
        )

        cell.value = header

    apply_header_style(sheet, 3)

    # ------------------------------------------------------
    # DATA
    # ------------------------------------------------------

    row = 4

    products = (
        Product.objects
        .select_related("category")
        .order_by("product_name")
    )

    total_stock = 0
    total_stock_value = Decimal("0.00")

    for product in products:

        stock_value = (
            product.current_stock *
            product.purchase_price
        )

        status = (
            "Low Stock"
            if product.current_stock <= 10
            else "Available"
        )

        sheet.cell(
            row=row,
            column=1,
            value=product.product_name,
        )

        sheet.cell(
            row=row,
            column=2,
            value=product.category.name,
        )

        sheet.cell(
            row=row,
            column=3,
            value=money(
                product.purchase_price
            ),
        )

        sheet.cell(
            row=row,
            column=4,
            value=money(
                product.default_selling_price
            ),
        )

        sheet.cell(
            row=row,
            column=5,
            value=product.current_stock,
        )

        sheet.cell(
            row=row,
            column=6,
            value=product.unit,
        )

        sheet.cell(
            row=row,
            column=7,
            value=money(stock_value),
        )

        sheet.cell(
            row=row,
            column=8,
            value=status,
        )

        total_stock += product.current_stock
        total_stock_value += stock_value

        row += 1

    # ------------------------------------------------------
    # TOTALS
    # ------------------------------------------------------

    sheet.cell(
        row=row + 1,
        column=6,
        value="TOTAL STOCK",
    ).font = BOLD_FONT

    sheet.cell(
        row=row + 1,
        column=7,
        value=total_stock,
    ).font = BOLD_FONT

    sheet.cell(
        row=row + 2,
        column=6,
        value="STOCK VALUE",
    ).font = BOLD_FONT

    sheet.cell(
        row=row + 2,
        column=7,
        value=money(total_stock_value),
    ).font = BOLD_FONT

    # ------------------------------------------------------
    # AUTO WIDTH
    # ------------------------------------------------------

    auto_width(sheet)

    return workbook_response(
        workbook,
        "Stock_Report.xlsx",
    )

# ==========================================================
# CUSTOMER LEDGER
# ==========================================================

def export_customer_ledger_excel():

    workbook = Workbook()

    sheet = workbook.active
    sheet.title = "Customer Ledger"

    # ------------------------------------------------------
    # TITLE
    # ------------------------------------------------------

    sheet["A1"] = "CUSTOMER LEDGER"
    sheet["A1"].font = TITLE_FONT

    headers = [

        "Customer",

        "Opening Balance",

        "Sales Amount",

        "Received Amount",

        "Outstanding Balance",

    ]

    for col, header in enumerate(headers, start=1):

        sheet.cell(
            row=3,
            column=col,
            value=header,
        )

    apply_header_style(sheet, 3)

    row = 4

    total_opening = Decimal("0.00")
    total_sales = Decimal("0.00")
    total_received = Decimal("0.00")
    total_balance = Decimal("0.00")

    customers = Customer.objects.order_by("name")

    for customer in customers:

        sales = (
            SalesInvoice.objects
            .filter(customer=customer)
            .aggregate(total=Sum("grand_total"))
            ["total"]
            or Decimal("0.00")
        )

        received = (
            CustomerPayment.objects
            .filter(customer=customer)
            .aggregate(total=Sum("amount"))
            ["total"]
            or Decimal("0.00")
        )

        balance = customer.current_balance

        sheet.cell(
            row=row,
            column=1,
            value=customer.name,
        )

        sheet.cell(
            row=row,
            column=2,
            value=money(customer.opening_balance),
        )

        sheet.cell(
            row=row,
            column=3,
            value=money(sales),
        )

        sheet.cell(
            row=row,
            column=4,
            value=money(received),
        )

        sheet.cell(
            row=row,
            column=5,
            value=money(balance),
        )

        total_opening += customer.opening_balance
        total_sales += sales
        total_received += received
        total_balance += balance

        row += 1

    # ------------------------------------------------------
    # TOTALS
    # ------------------------------------------------------

    sheet.cell(
        row=row + 1,
        column=1,
        value="TOTAL",
    ).font = BOLD_FONT

    sheet.cell(
        row=row + 1,
        column=2,
        value=money(total_opening),
    ).font = BOLD_FONT

    sheet.cell(
        row=row + 1,
        column=3,
        value=money(total_sales),
    ).font = BOLD_FONT

    sheet.cell(
        row=row + 1,
        column=4,
        value=money(total_received),
    ).font = BOLD_FONT

    sheet.cell(
        row=row + 1,
        column=5,
        value=money(total_balance),
    ).font = BOLD_FONT

    auto_width(sheet)

    return workbook_response(
        workbook,
        "Customer_Ledger.xlsx",
    )

# ==========================================================
# SUPPLIER LEDGER
# ==========================================================

def export_supplier_ledger_excel():

    workbook = Workbook()

    sheet = workbook.active
    sheet.title = "Supplier Ledger"

    # ------------------------------------------------------
    # TITLE
    # ------------------------------------------------------

    sheet["A1"] = "SUPPLIER LEDGER"
    sheet["A1"].font = TITLE_FONT

    # ------------------------------------------------------
    # HEADERS
    # ------------------------------------------------------

    headers = [

        "Supplier",

        "Opening Balance",

        "Purchase Amount",

        "Paid Amount",

        "Outstanding Balance",

    ]

    for col, header in enumerate(
        headers,
        start=1,
    ):

        sheet.cell(
            row=3,
            column=col,
            value=header,
        )

    apply_header_style(sheet, 3)

    # ------------------------------------------------------
    # DATA
    # ------------------------------------------------------

    row = 4

    total_opening = Decimal("0.00")
    total_purchase = Decimal("0.00")
    total_paid = Decimal("0.00")
    total_balance = Decimal("0.00")

    suppliers = Supplier.objects.order_by("name")

    for supplier in suppliers:

        purchase_amount = (
            PurchaseInvoice.objects
            .filter(supplier=supplier)
            .aggregate(total=Sum("grand_total"))
            .get("total")
            or Decimal("0.00")
        )

        paid_amount = (
            SupplierPayment.objects
            .filter(supplier=supplier)
            .aggregate(total=Sum("amount"))
            .get("total")
            or Decimal("0.00")
        )

        balance = supplier.current_balance

        sheet.cell(
            row=row,
            column=1,
            value=supplier.name,
        )

        sheet.cell(
            row=row,
            column=2,
            value=money(
                supplier.opening_balance
            ),
        )

        sheet.cell(
            row=row,
            column=3,
            value=money(
                purchase_amount
            ),
        )

        sheet.cell(
            row=row,
            column=4,
            value=money(
                paid_amount
            ),
        )

        sheet.cell(
            row=row,
            column=5,
            value=money(
                balance
            ),
        )

        total_opening += supplier.opening_balance
        total_purchase += purchase_amount
        total_paid += paid_amount
        total_balance += balance

        row += 1

    # ------------------------------------------------------
    # TOTALS
    # ------------------------------------------------------

    sheet.cell(
        row=row + 1,
        column=1,
        value="TOTAL",
    ).font = BOLD_FONT

    sheet.cell(
        row=row + 1,
        column=2,
        value=money(total_opening),
    ).font = BOLD_FONT

    sheet.cell(
        row=row + 1,
        column=3,
        value=money(total_purchase),
    ).font = BOLD_FONT

    sheet.cell(
        row=row + 1,
        column=4,
        value=money(total_paid),
    ).font = BOLD_FONT

    sheet.cell(
        row=row + 1,
        column=5,
        value=money(total_balance),
    ).font = BOLD_FONT

    # ------------------------------------------------------
    # AUTO WIDTH
    # ------------------------------------------------------

    auto_width(sheet)

    return workbook_response(
        workbook,
        "Supplier_Ledger.xlsx",
    )

# ==========================================================
# PROFIT REPORT
# ==========================================================

def export_profit_excel():

    workbook = Workbook()

    sheet = workbook.active
    sheet.title = "Profit Report"

    sheet["A1"] = "PROFIT REPORT"
    sheet["A1"].font = TITLE_FONT

    headers = [
        "Product",
        "Quantity Sold",
        "Sales Value",
        "Purchase Cost",
        "Profit",
    ]

    for col, header in enumerate(headers, start=1):
        sheet.cell(row=3, column=col, value=header)

    apply_header_style(sheet, 3)

    row = 4

    total_sales = Decimal("0.00")
    total_purchase = Decimal("0.00")
    total_profit = Decimal("0.00")

    products = Product.objects.order_by("product_name")

    for product in products:

        items = SalesInvoiceItem.objects.filter(
            product=product
        )

        qty = (
            items.aggregate(total=Sum("quantity"))
            .get("total")
            or 0
        )

        sales_value = (
            items.aggregate(total=Sum("total"))
            .get("total")
            or Decimal("0.00")
        )

        purchase_cost = (
            Decimal(qty) *
            product.purchase_price
        )

        profit = sales_value - purchase_cost

        sheet.append([
            product.product_name,
            qty,
            money(sales_value),
            money(purchase_cost),
            money(profit),
        ])

        total_sales += sales_value
        total_purchase += purchase_cost
        total_profit += profit

    sheet.append([])

    sheet.append([
        "TOTAL",
        "",
        money(total_sales),
        money(total_purchase),
        money(total_profit),
    ])

    sheet[row + len(products) + 1][0].font = BOLD_FONT

    auto_width(sheet)

    return workbook_response(
        workbook,
        "Profit_Report.xlsx",
    )


# ==========================================================
# CUSTOMER WISE SALES
# ==========================================================

def export_customer_sales_excel():

    workbook = Workbook()

    sheet = workbook.active

    sheet.title = "Customer Sales"

    sheet["A1"] = "CUSTOMER WISE SALES"
    sheet["A1"].font = TITLE_FONT

    headers = [
        "Customer",
        "Invoices",
        "Sales Amount",
    ]

    for col, header in enumerate(headers, start=1):
        sheet.cell(
            row=3,
            column=col,
            value=header,
        )

    apply_header_style(sheet, 3)

    row = 4

    customers = Customer.objects.order_by("name")

    for customer in customers:

        invoices = SalesInvoice.objects.filter(
            customer=customer
        )

        invoice_count = invoices.count()

        sales = (
            invoices.aggregate(
                total=Sum("grand_total")
            ).get("total")
            or Decimal("0.00")
        )

        sheet.append([
            customer.name,
            invoice_count,
            money(sales),
        ])

    auto_width(sheet)

    return workbook_response(
        workbook,
        "Customer_Sales.xlsx",
    )


# ==========================================================
# PRODUCT WISE SALES
# ==========================================================

def export_product_sales_excel():

    workbook = Workbook()

    sheet = workbook.active

    sheet.title = "Product Sales"

    sheet["A1"] = "PRODUCT WISE SALES"
    sheet["A1"].font = TITLE_FONT

    headers = [
        "Product",
        "Quantity Sold",
        "Sales Value",
    ]

    for col, header in enumerate(headers, start=1):
        sheet.cell(
            row=3,
            column=col,
            value=header,
        )

    apply_header_style(sheet, 3)

    products = Product.objects.order_by(
        "product_name"
    )

    for product in products:

        items = SalesInvoiceItem.objects.filter(
            product=product
        )

        quantity = (
            items.aggregate(
                total=Sum("quantity")
            ).get("total")
            or 0
        )

        sales = (
            items.aggregate(
                total=Sum("total")
            ).get("total")
            or Decimal("0.00")
        )

        sheet.append([
            product.product_name,
            quantity,
            money(sales),
        ])

    auto_width(sheet)

    return workbook_response(
        workbook,
        "Product_Sales.xlsx",
    )

